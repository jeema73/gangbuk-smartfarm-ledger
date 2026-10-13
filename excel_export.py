import io
from datetime import datetime
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import database

def create_styled_header(ws, row_idx: int, headers: List[str]):
    header_fill = PatternFill(start_color="2D5A43", end_color="2D5A43", fill_type="solid")
    header_font = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )
    align_center = Alignment(horizontal="center", vertical="center")

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = thin_border
    ws.row_dimensions[row_idx].height = 24

def auto_fit_columns(ws, min_width=12):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = max(max_len + 4, min_width)

def export_ledger_to_excel(orders: List[Dict[str, Any]], title_suffix: str = "") -> bytes:
    """납품 내역을 관공서 표준 납품대장 엑셀 파일 바이너리로 생성"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "납품대장"

    # 1. 상단 타이틀
    ws.merge_cells("A1:I1")
    title_cell = ws["A1"]
    title_text = "강북구 스마트팜 로컬푸드 납품대장"
    if title_suffix:
        title_text += f" ({title_suffix})"
    title_cell.value = title_text
    title_cell.font = Font(name="맑은 고딕", size=15, bold=True, color="2D5A43")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # 2. 헤더 행
    headers = ["연번", "주문일자", "납품처(매장명)", "작물품목", "단가(원/g)", "수량(g)", "금액(원)", "납품일자", "비고"]
    create_styled_header(ws, 3, headers)

    # 3. 데이터 입력
    thin_border = Border(
        left=Side(style='thin', color='E5DCCB'),
        right=Side(style='thin', color='E5DCCB'),
        top=Side(style='thin', color='E5DCCB'),
        bottom=Side(style='thin', color='E5DCCB')
    )
    regular_font = Font(name="맑은 고딕", size=10)
    bold_font = Font(name="맑은 고딕", size=10, bold=True)

    row_idx = 4
    serial_no = 1
    total_qty = 0.0
    total_amount = 0.0

    # 납품일자 오름차순으로 대장 정리
    sorted_orders = sorted(orders, key=lambda x: (x.get("delivery_date") or "", x.get("order_date") or ""))

    for ord_obj in sorted_orders:
        o_date = ord_obj.get("order_date", "")
        d_date = ord_obj.get("delivery_date") or o_date
        store = ord_obj.get("store_name", "")
        items = ord_obj.get("items", [])

        for it in items:
            crop = it.get("crop", "")
            unit_price = float(it.get("unit_price", 0))
            qty = float(it.get("qty", 0))
            amount = float(it.get("amount", round(qty * unit_price)))
            remarks = it.get("remarks", "")

            total_qty += qty
            total_amount += amount

            # 셀 채우기
            ws.cell(row=row_idx, column=1, value=serial_no).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=2, value=o_date).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=3, value=store).alignment = Alignment(horizontal="left")
            ws.cell(row=row_idx, column=4, value=crop).alignment = Alignment(horizontal="left")
            
            c_price = ws.cell(row=row_idx, column=5, value=unit_price)
            c_price.number_format = "#,##0.0"
            c_price.alignment = Alignment(horizontal="right")

            c_qty = ws.cell(row=row_idx, column=6, value=qty)
            c_qty.number_format = "#,##0"
            c_qty.alignment = Alignment(horizontal="right")

            c_amt = ws.cell(row=row_idx, column=7, value=amount)
            c_amt.number_format = "#,##0"
            c_amt.alignment = Alignment(horizontal="right")
            c_amt.font = bold_font

            ws.cell(row=row_idx, column=8, value=d_date).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=9, value=remarks).alignment = Alignment(horizontal="left")

            for col in range(1, 10):
                c = ws.cell(row=row_idx, column=col)
                c.border = thin_border
                if col != 7:
                    c.font = regular_font

            ws.row_dimensions[row_idx].height = 20
            row_idx += 1
            serial_no += 1

    # 4. 합계 행
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=5)
    sum_label = ws.cell(row=row_idx, column=1, value="총 합 계")
    sum_label.font = Font(name="맑은 고딕", size=10, bold=True, color="2D5A43")
    sum_label.alignment = Alignment(horizontal="center", vertical="center")
    sum_fill = PatternFill(start_color="F0F8F2", end_color="F0F8F2", fill_type="solid")

    for col in range(1, 10):
        c = ws.cell(row=row_idx, column=col)
        c.fill = sum_fill
        c.border = thin_border

    c_tot_qty = ws.cell(row=row_idx, column=6, value=total_qty)
    c_tot_qty.number_format = "#,##0"
    c_tot_qty.font = bold_font
    c_tot_qty.alignment = Alignment(horizontal="right")

    c_tot_amt = ws.cell(row=row_idx, column=7, value=total_amount)
    c_tot_amt.number_format = "#,##0"
    c_tot_amt.font = Font(name="맑은 고딕", size=11, bold=True, color="2D5A43")
    c_tot_amt.alignment = Alignment(horizontal="right")

    ws.row_dimensions[row_idx].height = 24

    auto_fit_columns(ws)

    # 5. 작물 단가 및 공급자 시트 덧붙이기
    ws_meta = wb.create_sheet(title="작물단가_공급자정보")
    headers_crops = ["작물명", "단가(1g)", "분류", "비고"]
    create_styled_header(ws_meta, 1, headers_crops)
    crops = database.get_all_crops()
    for c_idx, cr in enumerate(crops, start=2):
        ws_meta.cell(row=c_idx, column=1, value=cr.get("name", ""))
        ws_meta.cell(row=c_idx, column=2, value=cr.get("unit_price", 0))
        ws_meta.cell(row=c_idx, column=3, value=cr.get("category", ""))
        ws_meta.cell(row=c_idx, column=4, value=cr.get("remarks", ""))
    auto_fit_columns(ws_meta)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
